"""
Report 2: the exhaustive export -- every information_items row of one
analysis run, one per line, plus a Summary sheet. Moved here from the e2e
test script (2026-10-11) so the production lesson job builds the same file.
Still a minimal layout, not a finished Report 2 design (see
lomb_reporting_requirements_v1.md).
"""

import io


def build_report2_xlsx(items: list[dict]) -> bytes:
    """items: rows from StorageRepository.list_information_items(). Returns
    the .xlsx file as bytes (the caller stores it via a BlobStore)."""
    import openpyxl
    from openpyxl.styles import Font

    bold, plain = Font(name="Arial", bold=True), Font(name="Arial")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "All Items"
    headers = ["feature_key", "utterance_ref", "input_ref", "review_status", "gravity", "created_at", "output_json"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = bold
    for item in items:
        ws.append([item.get(h) for h in headers])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = plain
    for col_letter, width in zip("ABCDEFG", (20, 14, 50, 22, 9, 26, 60)):
        ws.column_dimensions[col_letter].width = width

    summary = wb.create_sheet("Summary")
    summary.append(["feature_key", "flagged_count"])
    for cell in summary[1]:
        cell.font = bold
    last_row = len(items) + 1
    for row_i, feature_key in enumerate(sorted({i["feature_key"] for i in items}), start=2):
        summary.cell(row=row_i, column=1, value=feature_key).font = plain
        summary.cell(row=row_i, column=2, value=f"=COUNTIF('All Items'!A2:A{last_row},A{row_i})").font = plain
    summary.column_dimensions["A"].width = 20
    summary.column_dimensions["B"].width = 14

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
